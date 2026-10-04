"""Build the Word report from docs/report/report.md, on top of the CCP template.

    python tools/build_report.py

The page setup and styles come from
docs/report/CT351_KRR_Smart_Parking_10-10_Report_Template.docx; its body is
replaced. The markdown source holds the prose. Lines in double braces are
directives that pull tables, traces and diagrams out of the running code,
so the report always shows what the engines actually do:

    {{titlepage}} {{rubric}} {{contributions}} {{toc}} {{lof}} {{lot}}
    {{entities}} {{frame Person}} {{frame Faculty Staff}} {{fol_rules R06 R10 ...}}
    {{fol_statements}} {{dl_roles}} {{dl_axioms sub|equiv|disjoint|restrict}} {{abox}}
    {{horn_table}} {{cycles}} {{fwd_table GOAL}} {{conflict_table}} {{bwd_table SCENARIO GOAL}}
    {{decisions}} {{test_table}} {{facts_compact}} {{prolog_rules}} {{all_firings}}
    {{include PATH}} {{tree}} {{screens}} {{participation}} {{compliance}} {{pagebreak}}
    {{stat NAME}}

"Table: caption" on the line before a table or table directive adds a numbered
caption. Images are written ![caption](path){h=CM}; without {h=CM} a PNG
with a dpi tag is placed at its printed size.

Run tools/finalize_report.py afterwards (Windows with Word) to fill in the
contents, figure and table lists and the page numbers.
"""

from __future__ import annotations

import re
import subprocess
import sys
from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
TEMPLATE = DOCS / "report" / "CT351_KRR_Smart_Parking_10-10_Report_Template.docx"
OUT = DOCS / "report" / "Smart_Parking_KRR_CCP_Report.docx"
sys.path.insert(0, str(ROOT))

from krr.frames.ontology import build_frames  # noqa: E402
from krr.knowledge_base.rules import CONSTRAINTS, RULES  # noqa: E402
from krr.knowledge_base.scenarios import SCENARIOS  # noqa: E402
from krr.logic.description_logic import (ROLE_FROM_PREDICATE, ROLES, TBOX,  # noqa: E402
                                         DLReasoner, build_abox)
from krr.logic.fol import check_statements  # noqa: E402
from krr.reasoning.backward import BackwardChainer  # noqa: E402
from krr.reasoning.forward import support  # noqa: E402
from krr.system import run  # noqa: E402

TEAM = [
    ("Syed Mohmmad Huzaifa", "AI-23022", "Problem, knowledge acquisition, frames and script",
     "Wrote the problem statement and the 13 domain requirements, ran the knowledge "
     "acquisition sessions, designed the 15 class frames and 38 instance frames with their "
     "facets and defaults, and wrote the ParkingSession script.",
     "Ch. 1, 2, 4; krr/frames/"),
    ("Maaz ur Rehman", "AI-23036", "Semantic network, description logic, trade-offs",
     "Built the semantic network generator with inheritance and path queries, wrote the DL "
     "TBox (T01-T28, D01-D17), the classifier and consistency checker, the OWL export and "
     "the expressivity and tractability analysis.",
     "Ch. 3, 5, 7, 12; krr/semantic_net/, krr/logic/description_logic.py"),
    ("Owais Tariq", "AI-23037", "FOL, Horn knowledge base, forward chaining",
     "Wrote the FOL vocabulary, the rule readings and statements F01-F12 with the model "
     "checker, converted the policy into Horn rules R01-R28 and IC1-IC7, and built the "
     "forward chainer with stratified negation and the Prolog export.",
     "Ch. 6, 8, 9.2-9.3; krr/logic/fol.py, krr/reasoning/forward.py"),
    ("Shaikh Muhammad Zain", "AI-23306", "Backward chaining, interfaces, testing, report",
     "Built the backward chainer, proof trees and why-not explainer, the CLI and Streamlit "
     "interface, the 12 test scenarios and 59 unit tests, and put the report together.",
     "Ch. 9.4-9.6, 10, 11; krr/reasoning/backward.py, app.py, tests/"),
]
RUBRIC = [
    ("Criterion 1: Problem Statement & Motivation",
     "Objective of the project and domain requirements are not clear or described.",
     "Some lack of clarity in objectives/purpose or domain justification.",
     "The project's objectives and domain motivation are clearly stated and justified."),
    ("Criterion 2: Structural KR (Frames & Semantic Networks)",
     "Model fails to represent necessary concepts, frames, slots, or semantic relationships.",
     "Model includes frames and semantic networks but lacks complete slots, facets, or "
     "default values.",
     "Model meets Proficient criteria with comprehensive, accurate frame structures and "
     "semantic networks."),
    ("Criterion 3: Formal Logic (FOL & Description Logic)",
     "Logical representations contain major syntax/semantic errors or are missing.",
     "FOL and DL formulas are present but contain minor syntax/quantifier errors.",
     "Model represents domain constraints perfectly using correct FOL formulas and "
     "Description Logic axioms."),
    ("Criterion 4: Inference Engine (Horn Clauses & Chaining)",
     "Rules are not converted to Horn clauses; failure to execute chaining.",
     "Rules are in Horn clauses, but Forward/Backward chaining traces have minor logical gaps.",
     "Horn clauses are correctly formulated with flawless, detailed Forward and Backward "
     "chaining traces."),
    ("Criterion 5: Team Participation & Individual Viva",
     "The individual did not contribute to the project and failed to answer viva questions.",
     "The individual contributed partially and demonstrated basic understanding during viva.",
     "The individual contributed significantly and demonstrated mastery of KR concepts "
     "during viva."),
]
COMPLIANCE = [
    ("Problem statement and motivation", "Ch. 1", 0),
    ("Domain requirements", "1.4, 2.5", 0),
    ("Limitations of traditional databases", "2.6", 0),
    ("Frames and scripts", "Ch. 4", 0),
    ("Semantic network", "Ch. 5", 1),
    ("FOL facts, rules and quantifiers", "Ch. 6", 2),
    ("Description logic axioms", "Ch. 7", 1),
    ("Expressivity vs tractability", "Ch. 12", 1),
    ("Horn clauses", "Ch. 8, App. B", 2),
    ("Forward chaining trace", "9.2-9.3, App. C", 2),
    ("Backward chaining trace", "9.4-9.5, App. D", 3),
    ("Working code", "Ch. 10, App. E", 3),
    ("Execution output", "10.8, App. F", 3),
    ("Testing", "Ch. 11", 3),
    ("Team participation statement", "App. G", 3),
    ("References", "References", 0),
]

MONO = "Consolas"
HEAD_FILL = "DCE6F1"
CODE_FILL = "F3F4F6"


# --------------------------------------------------------------------------------
# shared state
# --------------------------------------------------------------------------------

class Ctx:
    def __init__(self) -> None:
        self.main = run("morning_rush")
        self.fs = build_frames()
        self.figure = 0
        self.table = 0
        self.pending_caption: str | None = None
        self.stats = self._stats()

    def _stats(self) -> dict[str, str]:
        checks = sum(len(run(sc).check_expectations()) for sc in SCENARIOS.values())
        inferred = sum(1 for *_, inf in DLReasoner().classify() if inf)
        constants = len({t for f in self.main.forward.facts for t in f[1:]})
        try:
            out = subprocess.run([sys.executable, "-m", "unittest", "discover", "-s", "tests"],
                                 cwd=ROOT, capture_output=True, text=True)
            m = re.search(r"Ran (\d+) tests", out.stderr)
            unit = m.group(1) if m else "?"
        except OSError:
            unit = "?"
        kb = self.main.kb
        n_facts = sum(len(kb.facts_for(p)) for p in kb.predicates)
        return {"goal_checks": str(checks), "inferred": str(inferred), "constants": str(constants),
                "unit_tests": unit, "facts": str(n_facts),
                "derived": str(len(self.main.forward.firings)),
                "cycles": str(len(self.main.forward.cycles)),
                "sara_steps": f"{self.main.prove('allocate(sara, s_a1)').steps:,}"}


# --------------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------------

def shade(el, colour: str) -> None:
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), colour)
    pr.append(shd)


def add_field(paragraph, instruction: str, placeholder: str = " ") -> None:
    def fld(kind):
        r = OxmlElement("w:r")
        c = OxmlElement("w:fldChar")
        c.set(qn("w:fldCharType"), kind)
        r.append(c)
        return r
    paragraph._p.append(fld("begin"))
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instruction} "
    r.append(it)
    paragraph._p.append(r)
    paragraph._p.append(fld("separate"))
    paragraph.add_run(placeholder)
    paragraph._p.append(fld("end"))


def add_inline(paragraph, text: str, size: float | None = None, bold: bool = False) -> None:
    """Text with `code` spans in a monospace font and **bold** spans."""
    for part in re.split(r"(`[^`]*`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = MONO
            r.font.size = Pt((size or 11) - 1.5)
        elif part.startswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
            if size:
                r.font.size = Pt(size)
        else:
            r = paragraph.add_run(part)
            r.bold = bold
            if size:
                r.font.size = Pt(size)


def para(doc, text: str = "", size: float | None = None, align=None, bold=False,
         after: float | None = None, italic=False):
    p = doc.add_paragraph()
    if text:
        add_inline(p, text, size, bold)
    if italic:
        for r in p.runs:
            r.italic = True
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def code_block(doc, text: str, size: float = 7.5) -> None:
    size = max(size, 7.5)          # smaller code is hard to read on paper
    lines = text.rstrip("\n").split("\n")
    if max(len(line) for line in lines) > 104:
        # too wide for the page at this size: halve every run of 4+ spaces so
        # nested traces keep their shape without wrapping
        lines = [re.sub(r" {4,}", lambda m: " " * (len(m.group()) // 2), line)
                 for line in lines]
    p = doc.add_paragraph(style="Code")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.size = Pt(size)
        if i < len(lines) - 1:
            r.add_break()
    shade(p._p, CODE_FILL)


def caption(doc, ctx: Ctx, kind: str, text: str) -> None:
    if kind == "Figure":
        ctx.figure += 1
        n = ctx.figure
    else:
        ctx.table += 1
        n = ctx.table
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4 if kind == "Table" else 8)
    p.paragraph_format.keep_with_next = kind == "Table"
    p.add_run(f"{kind} ")
    add_field(p, f"SEQ {kind} \\* ARABIC", str(n))
    p.add_run(f": {text}")


def cell_margins(t, top: int, bottom: int) -> None:
    """Top and bottom cell padding in twips, so text never sits on a border."""
    tblpr = t._tbl.tblPr
    mar = tblpr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = OxmlElement("w:tblCellMar")
        tblpr.append(mar)
    for side, val in (("top", top), ("bottom", bottom)):
        el = mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            mar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")


def mark_size(p, pt: float) -> None:
    """Give the paragraph mark the table font size; an empty cell otherwise
    keeps the 11 pt body size and makes its row taller than the others."""
    ppr = p._p.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), str(round(pt * 2)))


def table(doc, ctx: Ctx, header: list[str], rows: list[list[str]],
          widths: list[float] | None = None, font: float = 8.5, mono_cols=()) -> None:
    # never below 8 pt for text and 7.5 pt for code, whatever the caller asks for
    font = max(font - 0.5, 8.0)
    mono = max(font - 1, 7.5)
    if ctx.pending_caption:
        caption(doc, ctx, "Table", ctx.pending_caption)
        ctx.pending_caption = None
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        shade(c._tc, HEAD_FILL)
    # repeat the header row on every page
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            if i in mono_cols:
                r = p.add_run(str(value))
                r.font.name = MONO
                r.font.size = Pt(mono)
            else:
                add_inline(p, str(value), font)
    cell_margins(t, top=20, bottom=20)
    # A short table moves to the next page whole; a long one may split, but the
    # header always stays with its first two rows.
    chars = sum(len(str(v)) for row in rows for v in row)
    whole = len(t.rows) <= 11 and chars <= 900
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = ri < len(t.rows) - 1 and (whole or ri < 2)
                mark_size(p, font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.8


def image(doc, ctx: Ctx, path: Path, text: str, max_h: float | None = None) -> None:
    """Insert a figure. With max_h (cm) the picture is scaled to that height;
    without it, a PNG that carries a dpi tag is placed at its printed size."""
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
        dpi = im.info.get("dpi", (0, 0))[0]
    if max_h is None and dpi:
        width = min(16.0, w / dpi * 2.54)
    else:
        width = min(16.0, (max_h or 11.0) * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width))
    caption(doc, ctx, "Figure", text)


# --------------------------------------------------------------------------------
# document setup
# --------------------------------------------------------------------------------

def setup() -> Document:
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.different_first_page_header_footer = True
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.space_after = Pt(4)
    for name, size, before in (("Heading 1", 15, 12), ("Heading 2", 12, 7),
                               ("Heading 3", 11, 5)):
        st = styles[name]
        st.font.size = Pt(size)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(4)
    cap = styles["Caption"]
    cap.font.size = Pt(9)
    cap.font.name = "Times New Roman"
    code = styles.add_style("Code", 1)
    code.font.name = MONO
    code.font.size = Pt(7.5)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0
    rpr = code.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), MONO)
    lb = styles["List Bullet"]
    lb.paragraph_format.space_after = Pt(2)
    lb.paragraph_format.line_spacing = 1.1
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)
    return doc


# --------------------------------------------------------------------------------
# front matter
# --------------------------------------------------------------------------------

def title_page(doc, ctx) -> None:
    def line(text, size, bold=True, after=6):
        p = para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=after)
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        return p

    line("NED UNIVERSITY OF ENGINEERING & TECHNOLOGY", 16, after=2)
    line("DEPARTMENT OF COMPUTER SCIENCE & INFORMATION TECHNOLOGY", 12, after=28)
    line("CT-351 KNOWLEDGE REPRESENTATION & REASONING", 13, after=2)
    line("COMPLEX COMPUTING PROBLEM (CCP)", 12, after=36)
    line("SMART PARKING ACCESS & SPACE ALLOCATION", 18, after=0)
    line("REASONING SYSTEM", 18, after=10)
    line("Project Report", 13, bold=False, after=40)
    line("Group Members", 12, after=6)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(("No.", "Name", "Roll No.")):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].add_run(h).bold = True
        shade(c._tc, HEAD_FILL)
    for i, (name, roll, *_) in enumerate(TEAM, 1):
        for c, v in zip(t.add_row().cells, (f"S{i}", name, roll)):
            c.text = v
    for row in t.rows:
        for c, w in zip(row.cells, (1.5, 8, 4)):
            c.width = Cm(w)
    para(doc, after=30)
    for label, value in (("Group Number", "12"), ("Semester", "7th"),
                         ("Course Instructor", "Miss Dure Shahwar")):
        line(f"{label}: {value}", 11, bold=False, after=8)
    para(doc, after=24)
    line(f"NED University of Engineering & Technology, {date.today():%B %Y}", 10,
         bold=False, after=0)
    page_break(doc)


def rubric_page(doc, ctx) -> None:
    p = para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    for i, text in enumerate(("Department of Computer Science and IT",
                              "Course Code: CT-351 | Course Title: Knowledge Representation "
                              "& Reasoning",
                              "Complex Computing Problem (CCP): Grading Rubric")):
        r = p.add_run(text)
        r.bold = True
        if i < 2:
            r.add_break()
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(("Student No.", "Name", "Roll No.")):
        t.rows[0].cells[i].text = h
        shade(t.rows[0].cells[i]._tc, HEAD_FILL)
    for i, (name, roll, *_) in enumerate(TEAM, 1):
        for c, v in zip(t.add_row().cells, (f"S{i}", name, roll)):
            c.text = v
    para(doc, after=4)
    header = ["Criteria and scales", "0 (Unsatisfactory)", "1 (Developing)", "2 (Proficient)",
              "S1", "S2", "S3", "S4"]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(8.5)
        shade(c._tc, HEAD_FILL)
    for row in RUBRIC + [("Total Marks (Out of 10):", "", "", "")]:
        cells = t.add_row().cells
        for i, v in enumerate(list(row) + ["", "", "", ""]):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(v)
            r.font.size = Pt(8.5)
            r.bold = i == 0
    for row in t.rows:
        for c, w in zip(row.cells, (3.4, 3.0, 3.0, 3.3, 0.9, 0.9, 0.9, 0.9)):
            c.width = Cm(w)
    para(doc, after=20)
    para(doc, "Teacher's Signature: ___________________________          "
              "Date: _______________")
    page_break(doc)


def contributions(doc, ctx) -> None:
    table(doc, ctx, ["Member", "Role / area", "Specific contribution", "Evidence"],
          [[f"{n} ({r})", area, contrib, ev] for n, r, area, contrib, ev in TEAM],
          [3.2, 3.0, 6.6, 3.7], font=8.5)


def toc_field(doc, ctx, instruction: str, placeholder: str) -> None:
    p = doc.add_paragraph()
    add_field(p, instruction, placeholder)


def end_section(doc, cols: int = 1, first: bool = False) -> None:
    """End the current section here with `cols` columns. The first section keeps
    the footer references and a blank first page (the title page)."""
    body_sect = doc.element.body.find(qn("w:sectPr"))
    sect = deepcopy(body_sect)
    for tag in ("w:type", "w:cols", "w:titlePg"):
        for el in sect.findall(qn(tag)):
            sect.remove(el)
    if not first:
        t = OxmlElement("w:type")
        t.set(qn("w:val"), "continuous")
        sect.insert(0, t)
        for tag in ("w:headerReference", "w:footerReference"):
            for el in sect.findall(qn(tag)):
                sect.remove(el)
    c = OxmlElement("w:cols")
    c.set(qn("w:num"), str(cols))
    c.set(qn("w:space"), "567")
    sect.find(qn("w:pgMar")).addnext(c)
    if first:
        c.addnext(OxmlElement("w:titlePg"))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p._p.get_or_add_pPr().append(sect)


def finish_sections(doc) -> None:
    """The body's own sectPr closes the last section: single column, numbered."""
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    # a trailing empty spacer after the last table would start a blank page
    while True:
        prev = sect.getprevious()
        if prev is None or prev.tag != qn("w:p") or "".join(prev.itertext()).strip():
            break
        if prev.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None:
            break
        body.remove(prev)
    # Word needs a paragraph after a closing table; make it 1pt high
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)
    tail.add_run("").font.size = Pt(1)
    for el in sect.findall(qn("w:titlePg")):
        sect.remove(el)
    t = OxmlElement("w:type")
    t.set(qn("w:val"), "continuous")
    sect.insert(0, t)


# --------------------------------------------------------------------------------
# generated tables
# --------------------------------------------------------------------------------

def frame_table(doc, ctx, names: list[str]) -> None:
    fs = ctx.fs
    rows = []
    for name in names:
        frame = fs.frames[name]
        own = frame.slots
        for slot, spec in fs.slot_specs(name).items():
            if slot not in own:
                continue            # inherited slots are shown in the parent's table
            source = name if slot in own else fs.defining_frame(name, slot, "type") or ""
            if slot not in own:
                origin = next((a for a in fs.ancestors(name) if slot in fs.frames[a].slots), "")
                source = f"inherited from {origin}"
            elif frame.isa and slot in fs.slot_specs(frame.isa):
                source = f"overrides {frame.isa}"
            else:
                source = "own slot"
            facets = [] if source == "own slot" else [source]
            if spec.range:
                facets.append("range {" + ", ".join(map(str, spec.range)) + "}")
            if spec.cardinality != "single":
                facets.append("multiple")
            if spec.if_needed is not None:
                facets.append("if-needed: " + (spec.if_needed.__doc__ or "").strip())
            label = f"{name}.{slot}" if len(names) > 1 else slot
            rows.append([label, spec.type, "" if spec.default is None else str(spec.default),
                         "; ".join(facets) or "-", spec.predicate or "-"])
    table(doc, ctx, ["Slot", "Type", "Default", "Other facets", "Predicate"], rows,
          [3.6, 2.8, 2.0, 6.0, 2.6], font=8)
    inst = [f.name for f in fs.instances() if f.isa in names]
    if inst:
        p = para(doc, "Instances: " + ", ".join(inst) + ".", size=10, after=6)
        p.paragraph_format.space_before = Pt(0)


def fol_rules(doc, ctx, ids: list[str]) -> None:
    by_id = {r.rid: r for r in RULES}
    rows = []
    for i, rid in enumerate(ids, 1):
        r = by_id[rid]
        rows.append([f"FOL-{i:02d} ({rid})", r.natural, r.to_fol(), r.explanation or "-"])
    table(doc, ctx, ["ID", "Natural-language rule", "FOL formula", "Explanation"], rows,
          [2.0, 4.0, 6.2, 4.3], font=8)


def fol_statements(doc, ctx) -> None:
    rows = [[st.sid, st.text, st.english, "holds" if ok else f"fails {ce}"]
            for st, ok, ce in ctx.main.fol_report()]
    table(doc, ctx, ["ID", "Formula", "Meaning", "Model check"], rows,
          [1.1, 7.3, 6.2, 1.9], font=8)


def dl_roles(doc, ctx) -> None:
    preds = {role: pred for pred, (role, _) in ROLE_FROM_PREDICATE.items()}
    table(doc, ctx, ["Role", "Domain → range", "Horn predicate"],
          [[r, d.split(" (")[0], preds.get(r, "")] for r, d in ROLES.items()],
          [3.2, 8.0, 4.0], font=8.5)


def dl_axioms(doc, ctx, kind: str) -> None:
    def is_restriction(ax):
        s = str(ax.right)
        return ax.kind == "subsumption" and ("∃" in s or "∀" in s)
    pick = {
        "sub": lambda ax: ax.kind == "subsumption" and not is_restriction(ax),
        "restrict": is_restriction,
        "equiv": lambda ax: ax.kind == "equivalence",
        "disjoint": lambda ax: ax.kind == "disjoint",
    }[kind]
    axioms = [ax for ax in TBOX if pick(ax)]
    if kind in ("sub", "disjoint"):
        # compact: several axioms per row (disjointness axioms are longer)
        per_row = 3 if kind == "sub" else 2
        cells = [f"{ax.aid}  {ax}" for ax in axioms]
        while len(cells) % per_row:
            cells.append("")
        rows = [cells[i:i + per_row] for i in range(0, len(cells), per_row)]
        table(doc, ctx, ["Axiom"] * per_row, rows, [16.5 / per_row] * per_row, font=8.5)
    else:
        table(doc, ctx, ["ID", "Axiom", "Meaning"], [[ax.aid, str(ax), ax.note] for ax in axioms],
              [1.2, 8.0, 7.3], font=8)


def abox(doc, ctx) -> None:
    ab = build_abox(ctx.main.forward.facts)
    lines = ab.assertions()
    concept = [a for a in lines if a.count(",") == 0][:8]
    role = [a for a in lines if a.count(",") == 1][:8]
    rows = [[c, r] for c, r in zip(concept, role)]
    table(doc, ctx, ["Concept assertions (sample)", "Role assertions (sample)"], rows,
          [8.2, 8.2], font=8, mono_cols=(0, 1))


def horn_table(doc, ctx) -> None:
    rows = [[r.rid, r.natural, r.text] for r in RULES]
    rows += [[c.cid, c.natural, c.text] for c in CONSTRAINTS]
    table(doc, ctx, ["ID", "Natural language", "Horn clause"], rows, [1.2, 6.3, 9.0],
          font=8, mono_cols=(2,))


def cycles(doc, ctx) -> None:
    rows = []
    for c in ctx.main.forward.cycles:
        fired = c.rules_fired
        rows.append([str(c.number), str(c.stratum), str(c.conflict_set), str(c.fired),
                     ", ".join(fired) if fired else "none (fixpoint)"])
    table(doc, ctx, ["Cycle", "Stratum", "Conflict set", "New facts", "Rules fired"], rows,
          [1.4, 1.6, 2.2, 2.0, 9.3], font=8.5)


def fwd_table(doc, ctx, goal: str) -> None:
    fwd = ctx.main.forward
    rows = []
    for f in support(fwd, goal):
        rows.append([str(f.step), f"{f.cycle} / {f.stratum}", " ∧ ".join(f.premises),
                     f.rule.rid, fmt(f.derived), f.rule.natural])
    table(doc, ctx, ["Step", "Cycle / stratum", "Facts before (premises)", "Rule", "New fact",
                     "Reason"], rows, [1.0, 1.4, 6.0, 1.0, 3.2, 3.9], font=7.5)


def conflict_table(doc, ctx) -> None:
    fwd = ctx.main.forward
    rows = []
    for f in fwd.firings:
        if f.rule.rid in ("R21", "R22", "R25", "R26", "R28") and \
                any(p.startswith("requests") or p.startswith("person(sara") or "s_a1" in p
                    for p in f.premises) or f.rule.rid == "R21":
            if f.rule.rid in ("R22",):
                continue
            b = {k: v for k, v in f.bindings.items() if k in f.rule.variables()}
            rows.append([str(f.step), f.rule.rid,
                         ", ".join(f"{k}={v}" for k, v in b.items()), fmt(f.derived)])
    table(doc, ctx, ["Step", "Rule", "Bindings", "New fact"], rows, [1.2, 1.2, 9.6, 4.5],
          font=8)


def fmt(atom) -> str:
    from krr.core.terms import format_atom
    return format_atom(atom)


def bwd_table(doc, ctx, scenario: str, goal: str) -> None:
    r = run(scenario)
    res = BackwardChainer(r.kb).prove(goal)
    rows = []

    def walk(node, depth):
        if node.how in ("fact", "built-in"):
            return
        indent = "· " * depth
        if node.how == "negation as failure":
            rows.append([indent + node.text, "NAF", "try to prove " + node.text[4:],
                         "no proof found", "success"])
            return
        subs = [c.text for c in node.children]
        evidence = [c.text for c in node.children if c.how in ("fact", "built-in")]
        rows.append([indent + node.text, node.how, ", ".join(subs),
                     ", ".join(evidence) or "(all derived)", "success"])
        for c in node.children:
            walk(c, depth + 1)

    walk(res.proofs[0], 0)
    table(doc, ctx, ["Goal", "Rule", "Subgoals", "Evidence (facts, built-ins)", "Result"], rows,
          [3.6, 1.0, 5.4, 5.0, 1.5], font=7.5)


def decisions(doc, ctx) -> None:
    rows = []
    for o in ctx.main.outcomes():
        rows.append([o.person, o.space, "allocated" if o.allocated else "denied",
                     "-" if o.allocated else "; ".join(o.reasons)])
    table(doc, ctx, ["Driver", "Space", "Decision", "Reason from the explainer"], rows,
          [1.6, 1.3, 1.8, 11.8], font=8)


TEST_INPUT = {
    "tc01_valid_student": "student(ali), permit p_ali valid to 20270630, s_a1 vacant",
    "tc02_expired_permit": "permit_expiry(p_usman, 20260915), today 20261002",
    "tc03_visitor_no_pass": "visitor(hina), no has_permit fact",
    "tc04_occupied": "requests(ali, s_a2), space_status(s_a2, occupied)",
    "tc05_student_in_employee_zone": "requests(ali, s_e1), zone_e is employee_zone",
    "tc06_no_suitable_space": "fatima owns motorbike, s_e1 is standard",
    "tc07_conflict": "zara 08:50, ali 09:05, sara 09:10 (badge) all request s_a1",
    "tc08_accessible_ineligible": "requests(ali, s_a3), accessibility_badge(ali, no)",
    "tc09_double_booking": "reserved_by(r2, ahmed), reserved_by(r3, ahmed), same date",
    "tc10_zone_closed": "current_hour(19), zone_v open 8-17",
    "tc11_reserved_by_other": "s_a5 reserved by zara (r1), ali requests it",
    "tc12_inconsistent_abox": "student(ali) plus faculty(ali)",
}


def test_table(doc, ctx) -> None:
    rows = []
    for key, sc in SCENARIOS.items():
        if key not in TEST_INPUT:
            continue
        checks = run(sc).check_expectations()
        focus = next((c for c in checks if c[0] == sc.focus_goal), checks[-1])
        goal, exp, fwd, bwd = focus
        ok = all(e == f == b for _, e, f, b in checks)
        tid, title = sc.title.split(None, 1)
        rows.append([tid.replace("TC", "TC-").replace("TC-", "TC-0") if len(tid) == 3 else
                     tid.replace("TC", "TC-"), title.strip(), TEST_INPUT[key],
                     f"{goal} is {str(exp).lower()}",
                     f"forward {str(fwd).lower()}, backward {str(bwd).lower()} "
                     f"({len(checks)} goals checked)", "PASS" if ok else "FAIL"])
    table(doc, ctx, ["TC", "Scenario", "Input facts", "Expected", "Actual", "Status"], rows,
          [1.2, 3.4, 4.0, 3.5, 3.4, 1.0], font=7.5)


def facts_compact(doc, ctx) -> None:
    text = (DOCS / "generated" / "facts.txt").read_text(encoding="utf-8")
    groups = re.findall(r"\[(.+?)\]\s+\d+ facts\n((?:    .+\n?)+)", text)
    for name, body in groups:
        facts = [f.strip().rstrip(".") for f in body.strip().split("\n")]
        p = para(doc, "", after=3)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(f"{name} ({len(facts)}): ")
        r.bold = True
        r.font.size = Pt(8.5)
        r = p.add_run(";  ".join(facts) + ".")
        r.font.name = MONO
        r.font.size = Pt(7.5)


def prolog_tail(doc, ctx) -> None:
    """The part of the generated Prolog file that differs from the Horn table:
    the negated decision rule, the constraints as violation/1 and the helpers."""
    text = (ROOT / "prolog" / "smart_parking.pl").read_text(encoding="utf-8")
    code_block(doc, text[text.index("% H. Decision"):], 7)


def prolog_rules(doc, ctx) -> None:
    lines = []
    for r in RULES:
        lines.append(r.text.replace("not ", "\\+ ") + f"  % {r.rid}")
    for c in CONSTRAINTS:
        lines.append(f"{c.text}  % {c.cid}")
    code_block(doc, "\n".join(lines), 6.5)


def all_firings(doc, ctx) -> None:
    items = [f"{f.step:>3} C{f.cycle} {f.rule.rid} {fmt(f.derived)}"
             for f in ctx.main.forward.firings]
    n = (len(items) + 1) // 2
    cols = [items[i * n:(i + 1) * n] for i in range(2)]
    width = max(len(s) for s in items) + 2
    lines = ["".join(c[i].ljust(width) if i < len(c) else "" for c in cols).rstrip()
             for i in range(n)]
    code_block(doc, "\n".join(lines), 6.5)


def screens(doc, ctx) -> None:
    """Cropped captures of the Streamlit app (1000 px viewport, 2x), one per row.
    Each PNG carries a dpi tag that makes it 15.5 cm wide, so the UI text
    prints at about 7 pt."""
    shots = [("report_overview", "Overview tab: decisions for the morning rush"),
             ("report_forward", "Forward chaining tab: cycle table"),
             ("report_proof_tree", "Test cases tab: proof tree for allocate(ali, s_a1)"),
             ("report_test_cases", "Test cases tab: both engines on all cases"),
             ("report_site_3d", "Project website: the 3D car park after the morning rush (intro panel hidden)"),
             ("report_site_why", "Project website: Ask why for omar, a driver added through the form")]
    for name, text in shots:
        image(doc, ctx, DOCS / "screenshots" / f"{name}.png", text)


def participation(doc, ctx) -> None:
    viva = {
        "AI-23022": "Problem, frames, script demo",
        "AI-23036": "Semantic network, DL reasoner demo",
        "AI-23037": "Horn rules, forward chaining demo",
        "AI-23306": "Backward chaining, UI, test cases demo",
    }
    rows = [[f"{n} ({r})", ev, viva[r]] for n, r, _, _, ev in TEAM]
    table(doc, ctx, ["Member", "Report sections and code", "Viva presentation part"], rows,
          [4.2, 7.8, 4.5], font=8.5)


def compliance(doc, ctx) -> None:
    rows = [[req, "Yes", where, TEAM[who][0]] for req, where, who in COMPLIANCE]
    table(doc, ctx, ["Requirement", "Completed?", "Evidence / section", "Verified by"], rows,
          [5.8, 2.0, 4.2, 4.5], font=8.5)


def tree() -> str:
    """Directories one per line, with their files listed after them."""
    skip = {"__pycache__", "png", "generated", "screenshots", "viva", "diagrams", "report"}
    lines = []

    def walk(path: Path, depth: int) -> None:
        files = sorted(p.name for p in path.iterdir() if p.is_file()
                       and not p.name.startswith(".") and p.name != "__init__.py"
                       and p.suffix not in (".pyc", ".mmd", ".pdf"))
        if not files and not any(path.iterdir()):
            return                  # empty folder
        if not files:               # list the sub-folders that are not expanded
            files = sorted(d.name + "/" for d in path.iterdir() if d.is_dir())
        name = "smart-parking-krr/" if depth == 0 else "  " * depth + path.name + "/"
        lines.append(f"{name:<24}{', '.join(files)}")
        if path.name == "docs":
            return
        for d in sorted(p for p in path.iterdir() if p.is_dir()):
            if d.name not in skip and not d.name.startswith("."):
                walk(d, depth + 1)

    walk(ROOT, 0)
    return "\n".join(lines)


def _old_tree() -> str:
    lines = ["smart-parking-krr/"]
    keep = {"krr", "core", "frames", "knowledge_base", "reasoning", "logic", "semantic_net",
            "export", "tests", "tools"}
    skip = {"__pycache__", ".git", "png", "generated", ".streamlit", "screenshots",
            "diagrams", "viva", ".pytest_cache", ".github", "report"}

    def walk(path: Path, prefix: str) -> None:
        entries = sorted([p for p in path.iterdir() if p.name not in skip
                          and not p.name.startswith(".") and p.suffix not in (".pyc", ".mmd")
                          and p.name != "__init__.py"],
                         key=lambda p: (p.is_file(), p.name))
        for i, e in enumerate(entries):
            last = i == len(entries) - 1
            lines.append(f"{prefix}{'└── ' if last else '├── '}{e.name}"
                         + ("/" if e.is_dir() else ""))
            if e.is_dir() and e.name in keep:
                walk(e, prefix + ("    " if last else "│   "))

    walk(ROOT, "")
    return "\n".join(lines)


ENTITIES = [
    ("Person", "Anyone who may drive onto campus", "fullName, accessibilityBadge, owns, hasPermit"),
    ("Student", "Enrolled student (Ali, Sara, Usman, Zara)", "studentId, department"),
    ("Faculty / Staff", "Employees: Dr. Ahmed (faculty), Fatima (staff)",
     "employeeId, department, designation / jobTitle"),
    ("Visitor", "Guest without a university permit (Bilal, Hina)", "purpose, hostDepartment"),
    ("Vehicle", "Registered car, electric car or motorbike", "plateNo, vehicleType"),
    ("Parking Permit", "Student permit, employee permit or one-day visitor pass",
     "permitClass, status, startDate, expiryDate"),
    ("Parking Space", "One bay; standard, accessible, EV charging or motorbike",
     "locatedIn, spaceType, status, zoneType"),
    ("Parking Zone", "Group of bays for one kind of user, with opening hours",
     "zoneType, openHour, closeHour, partOf, capacity"),
    ("Reservation", "A person's booking of one bay for one day", "madeBy, forSpace, date"),
]


# --------------------------------------------------------------------------------
# markdown rendering
# --------------------------------------------------------------------------------

def render(doc, ctx: Ctx, text: str) -> None:
    text = re.sub(r"\{\{stat (\w+)\}\}", lambda m: ctx.stats.get(m.group(1), "?"), text)
    lines = text.split("\n")
    i = 0
    buf: list[str] = []

    def flush():
        if buf:
            para(doc, " ".join(s.strip() for s in buf))
            buf.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        d = re.fullmatch(r"\{\{(.+?)\}\}", s)
        if d:
            flush()
            directive(doc, ctx, d.group(1).strip())
        elif s.startswith("Table: "):
            flush()
            ctx.pending_caption = s[7:]
        elif s.startswith("```"):
            flush()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            code_block(doc, "\n".join(block), 8)
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            title = s[level:].strip()
            if title.endswith("{newpage}"):
                title = title[:-9].strip()
                page_break(doc)
            doc.add_heading(title, level=level)
        elif s.startswith("!["):
            flush()
            m = re.match(r"!\[(.*?)\]\((.*?)\)(\{h=([\d.]+)\})?", s)
            image(doc, ctx, DOCS / m.group(2), m.group(1),
                  float(m.group(4)) if m.group(4) else None)
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, ctx, rows[0], rows[1:])
            continue
        elif s.startswith("- "):
            flush()
            add_inline(doc.add_paragraph(style="List Bullet"), s[2:])
        elif re.match(r"^\d+\. ", s):
            flush()
            num, rest = s.split(" ", 1)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Cm(0.8), Cm(-0.6)
            pf.space_after, pf.line_spacing = Pt(2), 1.1
            pf.tab_stops.add_tab_stop(Cm(0.8))
            add_inline(p, f"{num}\t{rest}")
        elif not s:
            flush()
        else:
            buf.append(line)
        i += 1
    flush()


def directive(doc, ctx: Ctx, d: str) -> None:
    name, _, arg = d.partition(" ")
    simple = {
        "titlepage": title_page, "rubric": rubric_page,
        "contributions": contributions, "fol_statements": fol_statements,
        "dl_roles": dl_roles, "abox": abox, "horn_table": horn_table, "cycles": cycles,
        "conflict_table": conflict_table, "decisions": decisions, "test_table": test_table,
        "facts_compact": facts_compact, "prolog_rules": prolog_rules,
        "prolog_tail": prolog_tail,
        "all_firings": all_firings, "screens": screens, "participation": participation,
        "compliance": compliance,
    }
    if name in simple:
        simple[name](doc, ctx)
    elif name == "toc":
        end_section(doc, first=True)
        toc_field(doc, ctx, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field.")
    elif name == "endcols":
        end_section(doc, cols=2)
    elif name == "lof":
        toc_field(doc, ctx, 'TOC \\h \\z \\c "Figure"', "Right-click and choose Update Field.")
    elif name == "lot":
        toc_field(doc, ctx, 'TOC \\h \\z \\c "Table"', "Right-click and choose Update Field.")
    elif name == "pagebreak":
        page_break(doc)
    elif name == "entities":
        table(doc, ctx, ["Entity", "Description", "Important attributes (frame slots)"],
              [list(e) for e in ENTITIES], [3.0, 6.5, 7.0], font=8.5)
    elif name == "frame":
        frame_table(doc, ctx, arg.split())
    elif name == "fol_rules":
        fol_rules(doc, ctx, arg.split())
    elif name == "dl_axioms":
        dl_axioms(doc, ctx, arg)
    elif name == "fwd_table":
        fwd_table(doc, ctx, arg)
    elif name == "bwd_table":
        scenario, goal = arg.split(" ", 1)
        bwd_table(doc, ctx, scenario, goal)
    elif name == "include":
        path, _, size = arg.partition(" ")
        code_block(doc, (DOCS / path).read_text(encoding="utf-8"), float(size or 7))
    elif name == "tree":
        code_block(doc, tree(), 7.5)
    else:
        raise ValueError(f"unknown directive {{{{{d}}}}}")


def build() -> Path:
    ctx = Ctx()
    doc = setup()
    render(doc, ctx, (DOCS / "report" / "report.md").read_text(encoding="utf-8"))
    finish_sections(doc)
    # document properties (Word copies title and author into the exported PDF)
    props = doc.core_properties
    props.title = "Smart Parking Access and Space Allocation Reasoning System"
    props.subject = "CT-351 Knowledge Representation and Reasoning, Complex Computing Problem"
    props.author = "Group 12: " + ", ".join(name for name, *_ in TEAM)
    props.last_modified_by = props.author
    props.comments = ""
    props.category = "CCP report"
    out = OUT
    try:
        doc.save(out)
    except PermissionError:
        out = out.with_name(out.stem + "_new.docx")
        doc.save(out)
        print("  the report is open in Word; saved a copy instead")
    print(f"  wrote {out.relative_to(ROOT)}")
    return out


if __name__ == "__main__":
    build()
